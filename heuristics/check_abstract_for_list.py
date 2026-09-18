"""
Etapa "Addition by Heuristics": dada uma lista de artigos candidatos a
falso negativo (ex.: títulos que o Filter by Title excluiu mas que você
quer reavaliar manualmente), verifica quais já têm abstract em algum
CSV intermediário gerado pelo pipeline e gera uma planilha com esses
artigos + o abstract quando existir, para revisão manual.

Busca em TODOS os arquivos intermediários que existirem em
data/processed/ (prioriza o pool bruto, que tem os abstracts originais
antes de qualquer filtro), casando por título normalizado — mais
robusto que URL, já que cada base usa um formato de link diferente.

Formato do arquivo de lista (uma linha por artigo, ver exemplos em
data/heuristic_lists/):
    source_db|~|título|~|ano|~|venue|~|url

Uso:
    python heuristics/check_abstract_for_list.py data/heuristic_lists/lista_1.txt
    python heuristics/check_abstract_for_list.py lista.txt saida.xlsx   # nome de saída opcional

Sem o nome de saída, gera data/processed/Abstracts_<nome_da_lista>.xlsx.
Para os artigos que ficarem SEM abstract, use em seguida
heuristics/fetch_abstracts_for_list.py (busca na OpenAlex / Semantic
Scholar) e heuristics/prioritize_list_review.py (ordena a leitura).
"""

import csv
import os
import re
import sys

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# Arquivos candidatos a conter abstract, em ordem de prioridade
# (o primeiro que tiver abstract não-vazio pra um título "ganha")
CANDIDATE_FILES = [
    os.path.join("data", "processed", "raw_pool_no_dedup.csv"),
    os.path.join("data", "processed", "after_language_filter.csv"),
    os.path.join("data", "processed", "after_title_filter_clean.csv"),
    os.path.join("data", "processed", "abstract_screening_results.csv"),
    os.path.join("data", "processed", "final_combined_deduplicated.csv"),
    os.path.join("data", "processed", "final_after_ec6.csv"),
    os.path.join("data", "processed", "combined_deduplicated.csv"),
]


def load_list(path: str) -> list[dict]:
    """Lê o arquivo de lista (formato source_db|~|título|~|ano|~|venue|~|url)."""
    items = []
    with open(path, encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line or line.startswith("#"):
                continue
            parts = line.split("|~|")
            if len(parts) != 5:
                print(f"AVISO: linha mal formatada, pulando: {line[:80]}")
                continue
            source_db, title, year, venue, url = parts
            items.append({"source_db": source_db, "title": title, "year": year,
                          "venue": venue, "url": url})
    return items


def normalize_title(title: str) -> str:
    t = title.lower()
    t = re.sub(r"[^a-z0-9\s]", " ", t)
    t = re.sub(r"\s+", " ", t).strip()
    return t


def build_abstract_index() -> dict:
    """Varre todos os arquivos candidatos e monta {título normalizado: abstract}."""
    index = {}
    files_found = []
    for path in CANDIDATE_FILES:
        if not os.path.exists(path):
            continue
        files_found.append(path)
        with open(path, encoding="utf-8") as f:
            reader = csv.DictReader(f)
            for row in reader:
                title_key = normalize_title(row.get("title", ""))
                abstract = row.get("abstract", "").strip()
                if title_key and abstract and title_key not in index:
                    index[title_key] = abstract

    print(f"Arquivos consultados: {len(files_found)}")
    for f in files_found:
        print(f"  {f}")
    print(f"Total de títulos com abstract disponível em algum arquivo: {len(index)}\n")
    return index


def main() -> None:
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    list_path = sys.argv[1]
    if len(sys.argv) > 2:
        out_path = sys.argv[2]
    else:
        list_name = os.path.splitext(os.path.basename(list_path))[0]
        out_path = os.path.join("data", "processed", f"Abstracts_{list_name}.xlsx")

    abstract_index = build_abstract_index()
    items = load_list(list_path)

    n_com_abstract = 0
    for item in items:
        title_key = normalize_title(item["title"])
        item["abstract"] = abstract_index.get(title_key, "")
        if item["abstract"]:
            n_com_abstract += 1

    print(f"Total de artigos na lista: {len(items)}")
    print(f"Com abstract encontrado: {n_com_abstract}")
    print(f"Sem abstract: {len(items) - n_com_abstract}")

    # Gera a planilha
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Lista com Abstracts"

    headers = ["#", "Base", "Título", "Ano", "Venue", "Link", "Tem Abstract?", "Abstract"]
    header_font = Font(name="Arial", bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2E5C8A", end_color="2E5C8A", fill_type="solid")
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = header_font
        c.fill = header_fill
        c.border = border
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for i, item in enumerate(items, start=1):
        r = i + 1
        tem_abstract = "Sim" if item["abstract"] else "Não"
        values = [i, item["source_db"], item["title"], item["year"],
                  item["venue"], item["url"], tem_abstract, item["abstract"]]
        for col, val in enumerate(values, start=1):
            cell = ws.cell(row=r, column=col, value=val)
            cell.font = Font(name="Arial")
            cell.border = border
            if col in (3, 8):
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            elif col == 6 and item["url"]:
                cell.hyperlink = item["url"]
                cell.font = Font(name="Arial", color="0563C1", underline="single")

    widths = [5, 15, 42, 7, 26, 32, 12, 60]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 30

    wb.save(out_path)
    print(f"\nPlanilha salva em {out_path}")


if __name__ == "__main__":
    main()